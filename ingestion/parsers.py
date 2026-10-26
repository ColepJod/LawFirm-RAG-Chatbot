"""Format-specific parsers that return LangChain Documents."""

from __future__ import annotations

from pathlib import Path

from langchain_core.documents import Document

from ingestion.hashing import sha256_file
from ingestion.metadata import base_metadata


def _read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="latin-1")


def _first_markdown_heading(text: str) -> str | None:
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("#"):
            return stripped.lstrip("#").strip() or None
    return None


def parse_text_file(
    path: Path,
    *,
    corpus_root: Path | None = None,
    document_hash: str | None = None,
) -> list[Document]:
    digest = document_hash or sha256_file(path)
    text = _read_text(path)
    metadata = base_metadata(
        path,
        corpus_root=corpus_root,
        document_hash=digest,
        page_number=None,
        section_heading=None,
    )
    return [Document(page_content=text, metadata=metadata)]


def parse_markdown_file(
    path: Path,
    *,
    corpus_root: Path | None = None,
    document_hash: str | None = None,
) -> list[Document]:
    digest = document_hash or sha256_file(path)
    text = _read_text(path)
    metadata = base_metadata(
        path,
        corpus_root=corpus_root,
        document_hash=digest,
        page_number=None,
        section_heading=_first_markdown_heading(text),
    )
    return [Document(page_content=text, metadata=metadata)]


def parse_pdf_file(
    path: Path,
    *,
    corpus_root: Path | None = None,
    document_hash: str | None = None,
    parser: str = "pdfmux",
) -> list[Document]:
    digest = document_hash or sha256_file(path)

    if parser == "pdfmux":
        from pdfmux.integrations.langchain import PDFMuxLoader

        raw_documents = PDFMuxLoader(path, quality="standard").load()
        documents: list[Document] = []
        for raw in raw_documents:
            raw_metadata = dict(getattr(raw, "metadata", {}) or {})
            page_number = raw_metadata.get("page_start")
            if not isinstance(page_number, int):
                page_number = None
            section_heading = raw_metadata.get("title")
            if not isinstance(section_heading, str) or not section_heading.strip():
                section_heading = None

            metadata = base_metadata(
                path,
                corpus_root=corpus_root,
                document_hash=digest,
                page_number=page_number,
                section_heading=section_heading,
            )
            merged_metadata = {**raw_metadata, **metadata}
            documents.append(Document(page_content=raw.page_content, metadata=merged_metadata))
        return documents

    if parser == "unstructured":
        from unstructured.partition.pdf import partition_pdf

        elements = partition_pdf(filename=str(path))
        text = "\n\n".join(str(element).strip() for element in elements if str(element).strip())
        metadata = base_metadata(
            path,
            corpus_root=corpus_root,
            document_hash=digest,
            page_number=None,
            section_heading=None,
        )
        return [Document(page_content=text, metadata=metadata)]

    if parser == "docling":
        from docling.document_converter import DocumentConverter

        converter = DocumentConverter()
        result = converter.convert(str(path))
        text = result.document.export_to_markdown()
        metadata = base_metadata(
            path,
            corpus_root=corpus_root,
            document_hash=digest,
            page_number=None,
            section_heading=None,
        )
        return [Document(page_content=text, metadata=metadata)]

    raise ValueError(f"Unsupported PDF parser: {parser}")


def _read_docx_with_python_docx(path: Path) -> str:
    from docx import Document as DocxDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = DocxDocument(str(path))
    blocks: list[str] = []

    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            if paragraph.text.strip():
                blocks.append(paragraph.text.strip())
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))

    return "\n\n".join(blocks)


def parse_docx_file(
    path: Path,
    *,
    corpus_root: Path | None = None,
    document_hash: str | None = None,
) -> list[Document]:
    digest = document_hash or sha256_file(path)

    try:
        from unstructured.partition.docx import partition_docx

        elements = partition_docx(filename=str(path))
        text = "\n\n".join(str(element).strip() for element in elements if str(element).strip())
    except ImportError:
        text = _read_docx_with_python_docx(path)

    metadata = base_metadata(
        path,
        corpus_root=corpus_root,
        document_hash=digest,
        page_number=None,
        section_heading=None,
    )
    return [Document(page_content=text, metadata=metadata)]


def parse_file(
    path: Path,
    *,
    corpus_root: Path | None = None,
    document_hash: str | None = None,
    pdf_parser: str = "pdfmux",
) -> list[Document]:
    """Parse one supported file into LangChain Documents."""

    path = Path(path)
    suffix = path.suffix.lower()
    digest = document_hash or sha256_file(path)

    if suffix in {".md", ".markdown"}:
        return parse_markdown_file(path, corpus_root=corpus_root, document_hash=digest)
    if suffix == ".txt":
        return parse_text_file(path, corpus_root=corpus_root, document_hash=digest)
    if suffix == ".pdf":
        return parse_pdf_file(
            path,
            corpus_root=corpus_root,
            document_hash=digest,
            parser=pdf_parser,
        )
    if suffix == ".docx":
        return parse_docx_file(path, corpus_root=corpus_root, document_hash=digest)

    raise ValueError(f"Unsupported file type: {path.suffix}")
